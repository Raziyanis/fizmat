"""Рабочая книга с самопроверкой: ученик пишет формулы, Excel сам проверяет ответ и способ записи."""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SC = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
FILLED = len(sys.argv) > 3 and sys.argv[3] == "filled"   # для теста: книга с правильными формулами

FN = {"СУММ": "SUM", "ABS": "ABS", "КОРЕНЬ": "SQRT", "СТЕПЕНЬ": "POWER", "СУММЕСЛИ": "SUMIF"}
def has_fn(t, ru):  # функция в тексте формулы: русское или английское имя
    return f'OR(ISNUMBER(SEARCH("{ru}(",{t})),ISNUMBER(SEARCH("{FN[ru]}(",{t})))'
def colons(t): return f'(LEN({t})-LEN(SUBSTITUTE({t},":","")))'

# Для каждого шага: ячейки-цели, ожидаемое значение (выражение от данных), условия на запись
STEPS = {
 "S1": [dict(c=["B8"], exp="SUM(B2:B7)", fn=["СУММ"], way=">=1:"),
        dict(c=["C8"], exp="SUM(C2:C7)", fn=["СУММ"]),
        dict(c=["E2"], exp="B2+B3+B6", fn=["СУММ"], way="0:"),
        dict(c=["E3"], exp="B2+B3+500", fn=["СУММ"], has=["500"]),
        dict(c=["E4"], exp="SUM(B2:B4)+B7", fn=["СУММ"], way="1:", has=["B7"]),
        dict(c=["E5"], exp="SUM(B2:B7)+SUM(C2:C7)", fn=["СУММ"], way="2:"),
        dict(c=["E6"], exp="SUM(B2:C7)", fn=["СУММ"], way="1:"),
        dict(c=["B8"], exp="SUM(B2:B7)", fn=["СУММ"], pre=("B7=1200", "сначала измени B7 на 1200"))],
 "S2": [dict(c=["G2"], exp="12", fn=["ABS"]),
        dict(c=["G3"], exp="ABS(B2)", fn=["ABS"]),
        dict(c=["E3", "E8"], exp=["C3-B3", "C8-B8"], nofn=["ABS"], fill="E2:E8"),
        dict(c=["D3", "D8"], exp=["ABS(C3-B3)", "ABS(C8-B8)"], fn=["ABS"], fill="D2:D8"),
        dict(c=["D4"], exp="ABS(C4-B4)", fn=["ABS"]),
        dict(c=["G4"], exp="SUMPRODUCT(ABS(C2:C8-B2:B8))", fn=["СУММ"]),
        dict(c=["G5"], exp="SUM(C2:C8)-SUM(B2:B8)", fn=["СУММ"]),
        dict(c=["G6"], exp="ABS(B8-B2)", fn=["ABS"])],
 "S3": [dict(c=["C2"], exp="SQRT(B2)", fn=["КОРЕНЬ"]),
        dict(c=["C5", "C6"], exp=["SQRT(B5)", "SQRT(B6)"], fill="C2:C6"),
        dict(c=["C4"], exp="SQRT(B4)"),
        dict(c=["F2"], exp="12", fn=["КОРЕНЬ"]),
        dict(c=["D2"], exp="SQRT(B2)*4"),
        dict(c=["F3"], exp="ERR", fn=["КОРЕНЬ"]),
        dict(c=["F4"], exp="5", fn=["КОРЕНЬ", "ABS"]),
        dict(c=["F5"], exp="SQRT(2*B2)", fn=["КОРЕНЬ"])],
 "S4": [dict(c=["C2"], exp="POWER(A2,B2)", fn=["СТЕПЕНЬ"]),
        dict(c=["C5"], exp="POWER(A5,B5)", fill="C2:C6"),
        dict(c=["C6"], exp="POWER(A6,B6)"),
        dict(c=["D3", "D6"], exp=["A3^B3", "A6^B6"], nofn=["СТЕПЕНЬ"], has=["^"], fill="D2:D6"),
        dict(c=["F2"], exp="POWER(A2,10)", fn=["СТЕПЕНЬ"], has=["10"]),
        dict(c=["F3"], exp="343", fn=["СТЕПЕНЬ"]),
        dict(c=["F4"], exp="9", fn=["СТЕПЕНЬ"], nofn=["КОРЕНЬ"]),
        dict(c=["F5"], exp="133100", fn=["СТЕПЕНЬ"])],
 "S5": [dict(c=["F2"], exp='SUMIF(C2:C9,">0")', fn=["СУММЕСЛИ"], has=[">0"]),
        dict(c=["F3"], exp='SUMIF(C2:C9,"<0")', fn=["СУММЕСЛИ"], has=["<0"]),
        dict(c=["F4"], exp='SUMIF(B2:B9,"еда",C2:C9)', fn=["СУММЕСЛИ"], has=["еда"]),
        dict(c=["F5"], exp="SUMIF(B2:B9,E5,C2:C9)", fn=["СУММЕСЛИ"], has=["E5"], pre=('E5<>""', "сначала впиши категорию в E5")),
        dict(c=["F6"], exp='SUMIF(C2:C9,"<-500")', fn=["СУММЕСЛИ"], has=["<-500"]),
        dict(c=["F7"], exp='-SUMIF(C2:C9,"<0")', fn=["СУММЕСЛИ", "ABS"]),
        dict(c=["F8"], exp="SUM(C2:C9)", fn=["СУММ"]),
        dict(c=["F5"], exp="SUMIF(B2:B9,E5,C2:C9)", fn=["СУММЕСЛИ"], pre=('E5="игры"', "сначала замени в E5 слово на «игры»"))],
 "S6": [dict(c=["E1"], exp="SQRT(B1)", fn=["КОРЕНЬ"]),
        dict(c=["E2"], exp="B1", fn=["СТЕПЕНЬ"]),
        dict(c=["E3"], exp="4*SQRT(B1)*B2"),
        dict(c=["E4"], exp="SUM(C6:C11)", fn=["СУММ"]),
        dict(c=["E5"], exp='SUMIF(B6:B11,"краска",C6:C11)', fn=["СУММЕСЛИ"]),
        dict(c=["E6"], exp="ABS(B3-SUM(C6:C11))", fn=["ABS"]),
        dict(c=["E7"], exp='SUM(C6:C11)-SUMIF(B6:B11,"краска",C6:C11)', fn=["СУММ", "СУММЕСЛИ"]),
        dict(c=["E6"], exp="ABS(B3-SUM(C6:C11))", fn=["ABS"], pre=("B3=60000", "сначала измени B3 на 60000"))],
}
# правильные формулы — только для тестовой «заполненной» книги
SOLVE = {
 "S1": {"B8": "=SUM(B2:B7)", "C8": "=SUM(C2:C7)", "E2": "=SUM(B2,B3,B6)", "E3": "=SUM(B2,B3,500)", "E4": "=SUM(B2:B4,B7)", "E5": "=SUM(B2:B7,C2:C7)", "E6": "=SUM(B2:C7)", "B7": 1200},
 "S2": {"G2": "=ABS(-12)", "G3": "=ABS(B2)", **{f"E{r}": f"=C{r}-B{r}" for r in range(2, 9)}, **{f"D{r}": f"=ABS(C{r}-B{r})" for r in range(2, 9)}, "G4": "=SUM(D2:D8)", "G5": "=SUM(E2:E8)", "G6": "=ABS(B8-B2)"},
 "S3": {**{f"C{r}": f"=SQRT(B{r})" for r in range(2, 7)}, "F2": "=SQRT(144)", "D2": "=C2*4", "F3": "=SQRT(-25)", "F4": "=SQRT(ABS(-25))", "F5": "=SQRT(C2*C2+C2*C2)"},
 "S4": {**{f"C{r}": f"=POWER(A{r},B{r})" for r in range(2, 7)}, **{f"D{r}": f"=A{r}^B{r}" for r in range(2, 7)}, "F2": "=POWER(A2,10)", "F3": "=POWER(7,3)", "F4": "=POWER(81,0.5)", "F5": "=100000*POWER(1.1,3)"},
 "S5": {"F2": '=SUMIF(C2:C9,">0")', "F3": '=SUMIF(C2:C9,"<0")', "F4": '=SUMIF(B2:B9,"еда",C2:C9)', "E5": "игры", "F5": "=SUMIF(B2:B9,E5,C2:C9)", "F6": '=SUMIF(C2:C9,"<-500")', "F7": '=ABS(SUMIF(C2:C9,"<0"))', "F8": "=SUM(C2:C9)"},
 "S6": {"E1": "=SQRT(B1)", "E2": "=POWER(E1,2)", "E3": "=4*E1*B2", "E4": "=SUM(C6:C11)", "E5": '=SUMIF(B6:B11,"краска",C6:C11)', "E6": "=ABS(B3-E4)", "E7": '=SUM(C6:C11)-SUMIF(B6:B11,"краска",C6:C11)', "B3": 60000},
}

def check_formula(st):
    cells = st["c"]; exps = st["exp"] if isinstance(st["exp"], list) else [st["exp"]]
    main = cells[0]; t = f'IFERROR(_xlfn.FORMULATEXT({main}),"")'
    conds = []  # (условие ошибки, текст)
    conds.append((" OR(" + ",".join(f"NOT(_xlfn.ISFORMULA({c}))" for c in cells) + ")", "✗ нужна формула, а не число" + (" (протяни формулу вниз)" if st.get("fill") else "")))
    for ru in st.get("fn", []): conds.append((f"NOT({has_fn(t, ru)})", f"✗ нужна функция {ru}"))
    for ru in st.get("nofn", []): conds.append((has_fn(t, ru), f"✗ здесь обойдись без {ru}"))
    way = []  # способ записи проверяем после ответа: сначала «верно ли посчитано», потом «тем ли способом»
    w = st.get("way")
    if w == "0:": way.append((f"{colons(t)}>0", "✗ ответ верный, но перечисли ячейки через «;», без диапазона"))
    if w == "1:": way.append((f"{colons(t)}<>1", "✗ ответ верный, но запиши одним диапазоном"))
    if w == "2:": way.append((f"{colons(t)}<>2", "✗ ответ верный, но нужны два диапазона через «;»"))
    if w == ">=1:": way.append((f"{colons(t)}<1", "✗ ответ верный, но используй диапазон через «:»"))
    for h in st.get("has", []):
        way.append((f'ISERROR(SEARCH("{h}",{t}))', "✗ ответ верный, но запиши так, как просит шаг"))
    if exps[0] == "ERR": ok = f"ISERROR({main})"
    else: ok = "AND(" + ",".join(f"IFERROR(ABS({c}-({e}))<=0.01*MAX(1,ABS({e})),FALSE)" for c, e in zip(cells, exps)) + ")"
    f = f'IF(ISBLANK({main}),"—",'
    closing = ")"
    if st.get("pre"):
        f += f'IF(NOT({st["pre"][0]}),"✗ {st["pre"][1]}",'; closing += ")"
    for c, msg in conds:
        f += f'IF({c},"{msg}",'; closing += ")"
    f += f'IF(NOT({ok}),"✗ ответ неверный — проверь формулу",'; closing += ")"
    for c, msg in way:
        f += f'IF({c},"{msg}",'; closing += ")"
    f += '"✓ верно"' + closing
    return "=" + f

thin = Side(style="thin", color="C9D3D9")
GREEN = PatternFill("solid", fgColor="D9F2E3"); RED = PatternFill("solid", fgColor="FBE1DE"); YEL = PatternFill("solid", fgColor="FFF4CC")
HEAD = PatternFill("solid", fgColor="1D7652"); WHITE = Font(bold=True, color="FFFFFF")
wb = Workbook(); info = wb.active; info.title = "Начало"
lines = [("Самопроверка: математические функции в Excel", True), ("", False),
         ("Класс:", True), ("Фамилия и имя:", True), ("", False),
         ("Как работать", True),
         ("1. Открой лист сценария (S1–S6). Таблица с данными уже заполнена.", False),
         ("2. Справа — шаги. Пиши формулу в жёлтую ячейку, которая указана в шаге.", False),
         ("3. В столбце «Проверка» Excel сразу покажет: ✓ верно — или что исправить.", False),
         ("4. Шаги со словами «измени…» — поменяй число в таблице, формулу не трогай.", False),
         ("5. Общий счёт — на листе «Итог».", False)]
for i, (txt, b) in enumerate(lines, 1):
    info.cell(i, 1, txt).font = Font(bold=b, size=14 if i == 1 else 11)
info["B3"].fill = YEL; info["B4"].fill = YEL
for k in ("B3", "B4"): info[k].protection = Protection(locked=False)
info.column_dimensions["A"].width = 22; info.column_dimensions["B"].width = 60
info.protection.sheet = True; info.protection.password = "fizmat"

summary = []
for s in SC:
    ws = wb.create_sheet(s["id"])
    for k, v in s["data"].items():
        t = v.replace(",", ".")
        ws[k] = int(t) if re.fullmatch(r"-?\d+", t) else (float(t) if re.fullmatch(r"-?\d+\.\d+", t) else v)
        if int(k[1:]) == 1 or v in ("Товар", "Категория", "Цена, тг"): ws[k].font = Font(bold=True)
    # всё на листе ученика можно править, кроме панели шагов
    for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=8):
        for c in row: c.protection = Protection(locked=False)
    targets = set()
    for st in STEPS[s["id"]]:
        if st.get("fill"):
            a, b = st["fill"].split(":"); col = a[0]
            for r in range(int(a[1:]), int(b[1:]) + 1): targets.add(f"{col}{r}")
        targets.update(st["c"])
    for k in targets: ws[k].fill = YEL
    c0 = 10  # J
    for j, h in enumerate(["№", "Тип шага", "Задание", "Ячейка", "Проверка"]):
        c = ws.cell(1, c0 + j, h); c.font = WHITE; c.fill = HEAD
    for i, (t, st) in enumerate(zip(s["tasks"], STEPS[s["id"]]), 1):
        r = i + 1
        cell_txt = (f"{t['cell']} → протяни вниз; ответ в {t['read']}" if t.get("read") else t["cell"])
        vals = [i, t["tag"], t["q"], cell_txt]
        for j, v in enumerate(vals):
            c = ws.cell(r, c0 + j, v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
        chk = ws.cell(r, c0 + 4, check_formula(st))
        chk.alignment = Alignment(wrap_text=True, vertical="top"); chk.border = Border(bottom=thin)
        chk.font = Font(bold=True)
        chk.protection = Protection(locked=True, hidden=True)   # формулу проверки не видно
    n = len(s["tasks"])
    rng = f"N2:N{n + 1}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(N2,1)="✓"'], fill=GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(N2,1)="✗"'], fill=RED))
    ws.cell(n + 3, c0 + 2, "Выполнено шагов:").font = Font(bold=True)
    sc_cell = ws.cell(n + 3, c0 + 4, f'=COUNTIF({rng},"✓*")&" из {n}"'); sc_cell.font = Font(bold=True, size=13)
    sc_cell.protection = Protection(locked=True, hidden=True)
    for col, w in {c0: 4, c0 + 1: 20, c0 + 2: 60, c0 + 3: 16, c0 + 4: 34}.items(): ws.column_dimensions[get_column_letter(col)].width = w
    for col in "ABCDEFGH": ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 18
    ws.freeze_panes = "A2"
    ws.protection.sheet = True; ws.protection.password = "fizmat"
    ws.protection.formatCells = False; ws.protection.formatColumns = False  # можно менять разрядность и ширину столбцов
    summary.append((s["id"], s["title"], s["fn"], rng, n))
    if FILLED:
        for k, v in SOLVE[s["id"]].items(): ws[k] = v

it = wb.create_sheet("Итог")
it["A1"] = "Итог самопроверки"; it["A1"].font = Font(bold=True, size=14)
it["A2"] = '="Ученик: "&Начало!B4&", "&Начало!B3'
for j, h in enumerate(["Сценарий", "Название", "Функция", "Верно", "Всего"]):
    c = it.cell(4, 1 + j, h); c.font = WHITE; c.fill = HEAD
for i, (sid, title, fn, rng, n) in enumerate(summary, 5):
    it.cell(i, 1, sid); it.cell(i, 2, title); it.cell(i, 3, fn)
    it.cell(i, 4, f"=COUNTIF({sid}!{rng},\"✓*\")"); it.cell(i, 5, n)
last = 4 + len(summary)
it.cell(last + 1, 2, "Всего").font = Font(bold=True)
it.cell(last + 1, 4, f"=SUM(D5:D{last})").font = Font(bold=True)
it.cell(last + 1, 5, f"=SUM(E5:E{last})").font = Font(bold=True)
it.column_dimensions["B"].width = 34; it.column_dimensions["C"].width = 16
it.protection.sheet = True; it.protection.password = "fizmat"
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel пересчитает проверку при открытии
wb.save(OUT)
print("saved", OUT)
