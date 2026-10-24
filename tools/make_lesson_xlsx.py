"""Этап «Работа в Excel»: файл самопроверки для учеников и ключ для учителя (уроки «Логические функции» и «Статистические функции»).

Сценарии берутся из index.html урока (EX_SC), поэтому условия на странице и проверки в файле совпадают.
Запуск: python3 tools/make_lesson_xlsx.py logic|stats [kz] [filled|wrong <файл>]  — без filled/wrong делает оба файла урока;
kz — казахская версия (тексты и данные из казахской части страницы, имена функций как в русском Excel);
filled/wrong — тестовая книга с правильными или с ошибочными формулами (для проверки самой проверки).
"""
import json, re, subprocess, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent.parent
LESSONS = {
    "logic": dict(dir="lessons-8-logic", check="Самопроверка_логика.xlsx", key="Ключ_логика.xlsx", topic="логические функции", title="Логические функции", fns="ЕСЛИ, И, ИЛИ, НЕ, СУММ",
                  kz=dict(check="Өзін-өзі_тексеру_логика.xlsx", key="Кілт_логика.xlsx", topic="логикалық функциялар", title="Логикалық функциялар")),
    "stats": dict(dir="lessons-8-stats", check="Самопроверка_статистика.xlsx", key="Ключ_статистика.xlsx", topic="статистические функции", title="Статистические функции", fns="СРЗНАЧ, МАКС, МИН, СУММ, ЕСЛИ",
                  kz=dict(check="Өзін-өзі_тексеру_статистика.xlsx", key="Кілт_статистика.xlsx", topic="статистикалық функциялар", title="Статистикалық функциялар")),
}
ARGS = sys.argv[1:]
KZ = "kz" in ARGS
if KZ: ARGS.remove("kz")
CFG = LESSONS[ARGS[0]]
if KZ: CFG = {**CFG, **CFG["kz"]}
LESSON = ROOT / CFG["dir"]
MODE = ARGS[1] if len(ARGS) > 1 else ""
def tr(ru, kz): return kz if KZ else ru

# Сценарии: русские — из src-ru.html, казахские — из казахской части собранной страницы (function L_kz)
js = ('const s0=require("fs").readFileSync(process.argv[1],"utf8"),s=s0.slice(s0.indexOf(process.argv[2]));'
      'eval(s.slice(s.indexOf("const EX_SC = ["),s.indexOf("EX_SC.forEach")).replace("const EX_SC","var EX_SC"));'
      'console.log(JSON.stringify(EX_SC))')
SC = json.loads(subprocess.run(["node", "-e", js, str(LESSON / ("index.html" if KZ else "src-ru.html")), "function L_kz()" if KZ else ""],
                               capture_output=True, text=True, check=True).stdout)

EN = {"СУММЕСЛИ": "SUMIF", "СУММ": "SUM", "СРЗНАЧ": "AVERAGE", "МАКС": "MAX", "МИН": "MIN", "ЕСЛИ": "IF", "ИЛИ": "OR", "И": "AND", "НЕ": "NOT"}
def outside_quotes(f, fn):
    return "".join(p if p.startswith('"') else fn(p) for p in re.split(r'("[^"]*")', f))
def to_en(f):
    """Русская запись Excel → формула для файла: имена функций, «;» → «,», десятичная запятая → точка."""
    def conv(p):
        p = re.sub(r"(\d),(\d)", r"\1.\2", p)
        p = re.sub(r"(?<![А-ЯЁA-Z])(СУММЕСЛИ|СУММ|СРЗНАЧ|МАКС|МИН|ЕСЛИ|ИЛИ|И|НЕ)\(", lambda m: EN[m.group(1)] + "(", p)
        return p.replace(";", ",")
    return outside_quotes(f, conv)
def shift(f, dr):
    return outside_quotes(f, lambda p: re.sub(r"(?<![A-ZА-ЯЁ])([A-I])(\d+)", lambda m: m.group(1) + str(int(m.group(2)) + dr), p))
def fill_cells(rng):
    a, b = rng.split(":"); return [f"{a[0]}{r}" for r in range(int(a[1:]), int(b[1:]) + 1)]
def num(v):
    if v == "": return None   # «очисти ячейку»
    t = v.replace(",", ".")
    return int(t) if re.fullmatch(r"-?\d+", t) else float(t) if re.fullmatch(r"-?\d+\.\d+", t) else v
def lit(v):
    return '""' if v == "" else str(num(v)) if not isinstance(num(v), str) else f'"{v}"'

def has_fn(t, ru):
    """Есть ли функция в тексте формулы (русское или английское имя). «И(» не путаем с ЕСЛИ( и ИЛИ(."""
    if ru == "И":
        return (f'OR(ISNUMBER(SEARCH("AND(",{t})),(LEN({t})-LEN(SUBSTITUTE(UPPER({t}),"И(","")))/2'
                f'>(LEN({t})-LEN(SUBSTITUTE(UPPER({t}),"ЛИ(","")))/3)')
    return f'OR(ISNUMBER(SEARCH("{ru}(",{t})),ISNUMBER(SEARCH("{EN[ru]}(",{t})))'

def steps(sc):
    """Для каждого шага: ячейки, ожидаемые выражения (живой расчёт по данным листа), функции, предусловие."""
    out = []
    for t in sc["tasks"]:
        cells = fill_cells(t["fill"]) if t.get("fill") else [t["cell"]]
        first = int(cells[0][1:]) if t.get("fill") else int(t["cell"][1:])
        base = int(t["fill"].split(":")[0][1:]) if t.get("fill") else first
        exps = [to_en(shift(t["f"], int(c[1:]) - base))[1:] for c in cells]
        pre = None
        if t.get("change"):
            cond = ",".join(f"{k}={lit(v)}" for k, v in t["change"].items())
            pre = (f"AND({cond})", tr("сначала ", "алдымен ") + ", ".join(clear_or_change(k, v) for k, v in t["change"].items()))
        out.append(dict(c=cells, exp=exps, fn=t.get("fn", []), num=isinstance(t["ans"], (int, float)), pre=pre, fill=bool(t.get("fill"))))
    return out

def clear_or_change(k, v, arrow=False):
    if v == "": return tr(f"очисти {k}", f"{k} ұяшығын тазала")
    return tr(f"измени {k} → {v}" if arrow else f"измени {k} на {v}", f"{k} мәнін {v} деп өзгерт")

def check_formula(st):
    cells, exps = st["c"], st["exp"]
    main = cells[0]; t = f'IFERROR(_xlfn.FORMULATEXT({main}),"")'
    f = f'IF(ISBLANK({main}),"—",'; closing = ")"
    if st["pre"]:
        f += f'IF(NOT({st["pre"][0]}),"✗ {st["pre"][1]}",'; closing += ")"
    f += 'IF(OR(' + ",".join(f"NOT(_xlfn.ISFORMULA({c}))" for c in cells) + '),"' + tr("✗ нужна формула, а не готовый ответ", "✗ дайын жауап емес, формула керек") + (tr(" (протяни формулу вниз)", " (формуланы төмен созып көшір)") if st["fill"] else "") + '",'; closing += ")"
    for ru in st["fn"]:
        f += f'IF(NOT({has_fn(t, ru)}),"{tr("✗ нужна функция " + ru, "✗ " + ru + " функциясы керек")}",'; closing += ")"
    if st["num"]:
        ok = "AND(" + ",".join(f"IFERROR(ABS({c}-({e}))<=0.01*MAX(1,ABS({e})),FALSE)" for c, e in zip(cells, exps)) + ")"
    else:
        norm = lambda x: f'SUBSTITUTE(LOWER({x}&""),"ё","е")'
        ok = "AND(" + ",".join(f"IFERROR({norm(c)}={norm('(' + e + ')')},FALSE)" for c, e in zip(cells, exps)) + ")"
    f += f'IF(NOT({ok}),"{tr("✗ ответ неверный — проверь формулу и диапазон", "✗ жауап дұрыс емес — формула мен ауқымды тексер")}",'; closing += ")"
    return "=" + f + '"' + tr("✓ верно", "✓ дұрыс") + '"' + closing

thin = Side(style="thin", color="C9D3D9")
GREEN = PatternFill("solid", fgColor="D9F2E3"); RED = PatternFill("solid", fgColor="FBE1DE"); YEL = PatternFill("solid", fgColor="FFF4CC")
HEAD = PatternFill("solid", fgColor="1D7652"); WHITE = Font(bold=True, color="FFFFFF")
C0 = 11   # панель шагов начинается со столбца K: данные занимают A–I

def put_data(ws, sc):
    for k, v in sc["data"].items():
        ws[k] = num(v)
        if int(k[1:]) == 1: ws[k].font = Font(bold=True)
    for col in "ABCDEFGHI": ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "A2"

def solve(ws, sc, wrong=False, changes=True):
    """Правильные формулы (и изменения данных) — для тестовой книги и ключа."""
    for t in sc["tasks"]:
        if t.get("change"):
            if changes:
                for k, v in t["change"].items(): ws[k] = num(v)
            continue
        cells = fill_cells(t["fill"]) if t.get("fill") else [t["cell"]]
        base = int(cells[0][1:])
        for c in cells:
            f = shift(t["f"], int(c[1:]) - base)
            if wrong: f = f.replace(">=", ">").replace('"зачёт";"незачёт"', '"незачёт";"зачёт"').replace('"есептелді";"есептелмеді"', '"есептелмеді";"есептелді"').replace("МИН(", "МАКС(").replace("СРЗНАЧ(", "СУММ(")
            ws[c] = to_en(f)

def selfcheck(path, mode=""):
    START, TOTAL = tr("Начало", "Басы"), tr("Итог", "Қорытынды")
    wb = Workbook(); info = wb.active; info.title = START
    lines = [(tr(f"Самопроверка: {CFG['topic']} в Excel", f"Өзін-өзі тексеру: Excel-дегі {CFG['topic']}"), True), ("", False),
             (tr("Класс:", "Сынып:"), True), (tr("Фамилия и имя:", "Тегі мен аты:"), True), ("", False),
             (tr("Как работать", "Қалай жұмыс істеу керек"), True)] + [(x, False) for x in tr([
             "1. Открой лист сценария (S1–S5). Таблица с данными уже заполнена.",
             "2. Справа — шаги. Пиши формулу в жёлтую ячейку, которая указана в шаге.",
             "3. «Протяни вниз» — потяни формулу за маленький квадрат в правом нижнем углу ячейки.",
             "4. В столбце «Проверка» Excel сразу покажет: ✓ верно — или что исправить.",
             "5. Шаги «измени…» — поменяй данные в таблице, формулы не трогай.",
             "6. Общий счёт — на листе «Итог»."], [
             "1. Сценарий парағын аш (S1–S5). Деректері бар кесте дайын.",
             "2. Оң жақта — қадамдар. Формуланы қадамда көрсетілген сары ұяшыққа жаз.",
             "3. «Төмен созып көшір» — формуланы ұяшықтың оң жақ төменгі бұрышындағы кішкентай шаршыдан тарт.",
             "4. «Тексеру» бағанында Excel бірден көрсетеді: ✓ дұрыс — немесе нені түзету керек.",
             "5. «Өзгерт…» қадамдары — кестедегі деректерді өзгерт, формулаларға тиме.",
             "6. Жалпы есеп — «Қорытынды» парағында."])]
    for i, (txt, b) in enumerate(lines, 1): info.cell(i, 1, txt).font = Font(bold=b, size=14 if i == 1 else 11)
    for k in ("B3", "B4"): info[k].fill = YEL; info[k].protection = Protection(locked=False)
    info.column_dimensions["A"].width = 22; info.column_dimensions["B"].width = 70
    info.protection.sheet = True; info.protection.password = "fizmat"
    summary = []
    for sc in SC:
        ws = wb.create_sheet(sc["id"]); put_data(ws, sc)
        for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=9):
            for c in row: c.protection = Protection(locked=False)
        sts = steps(sc)
        for st in sts:
            for k in st["c"]: ws[k].fill = YEL
        for j, h in enumerate(tr(["№", "Тип шага", "Задание", "Ячейка", "Проверка"], ["№", "Қадам түрі", "Тапсырма", "Ұяшық", "Тексеру"])):
            c = ws.cell(1, C0 + j, h); c.font = WHITE; c.fill = HEAD
        for i, (t, st) in enumerate(zip(sc["tasks"], sts), 1):
            r = i + 1
            a, b = t["fill"].split(":") if t.get("fill") else ("", "")
            ans = tr(f"; ответ в {t['cell']}", f"; жауап {t['cell']} ұяшығында")
            where = (tr(f"{a} → протяни до {b}", f"{a} → {b} ұяшығына дейін созып көшір") + ans if t.get("fill")
                     else ", ".join(clear_or_change(k, v, True) for k, v in t["change"].items()) + ans if t.get("change") else t["cell"])
            for j, v in enumerate([i, t["tag"], t["q"], where]):
                c = ws.cell(r, C0 + j, v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
            chk = ws.cell(r, C0 + 4, check_formula(st))
            chk.alignment = Alignment(wrap_text=True, vertical="top"); chk.border = Border(bottom=thin); chk.font = Font(bold=True)
            chk.protection = Protection(locked=True, hidden=True)   # формулу проверки не видно
        n = len(sc["tasks"]); col = get_column_letter(C0 + 4); rng = f"{col}2:{col}{n + 1}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}2,1)="✓"'], fill=GREEN))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}2,1)="✗"'], fill=RED))
        ws.cell(n + 3, C0 + 2, tr("Выполнено шагов:", "Орындалған қадамдар:")).font = Font(bold=True)
        sc_cell = ws.cell(n + 3, C0 + 4, f'=COUNTIF({rng},"✓*")&"{tr(" из ", " / ")}{n}"'); sc_cell.font = Font(bold=True, size=13)
        sc_cell.protection = Protection(locked=True, hidden=True)
        for c, w in {C0: 4, C0 + 1: 20, C0 + 2: 60, C0 + 3: 22, C0 + 4: 34}.items(): ws.column_dimensions[get_column_letter(c)].width = w
        ws.protection.sheet = True; ws.protection.password = "fizmat"
        ws.protection.formatCells = False; ws.protection.formatColumns = False
        summary.append((sc["id"], sc["title"], sc["fn"], rng, n))
        if mode: solve(ws, sc, wrong=mode == "wrong")
    it = wb.create_sheet(TOTAL)
    it["A1"] = tr("Итог самопроверки", "Өзін-өзі тексеру қорытындысы"); it["A1"].font = Font(bold=True, size=14)
    it["A2"] = f'="{tr("Ученик", "Оқушы")}: "&{START}!B4&", "&{START}!B3'
    for j, h in enumerate(tr(["Сценарий", "Название", "Функции", "Верно", "Всего"], ["Сценарий", "Атауы", "Функциялар", "Дұрыс", "Барлығы"])):
        c = it.cell(4, 1 + j, h); c.font = WHITE; c.fill = HEAD
    for i, (sid, title, fn, rng, n) in enumerate(summary, 5):
        it.cell(i, 1, sid); it.cell(i, 2, title); it.cell(i, 3, fn)
        it.cell(i, 4, f'=COUNTIF({sid}!{rng},"✓*")'); it.cell(i, 5, n)
    last = 4 + len(summary)
    it.cell(last + 1, 2, tr("Всего", "Барлығы")).font = Font(bold=True)
    it.cell(last + 1, 4, f"=SUM(D5:D{last})").font = Font(bold=True)
    it.cell(last + 1, 5, f"=SUM(E5:E{last})").font = Font(bold=True)
    it.column_dimensions["B"].width = 26; it.column_dimensions["C"].width = 22
    it.protection.sheet = True; it.protection.password = "fizmat"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel пересчитает проверку при открытии
    wb.save(path); print("saved", path)

def key(path):
    wb = Workbook(); info = wb.active; info.title = tr("Как пользоваться", "Қалай қолдану")
    for i, r in enumerate(tr([f"Ключ к этапу «Работа в Excel» (урок «{CFG['title']}», 8 класс).",
                           "На листах S1–S5 — таблица сценария и все формулы учеников уже вписаны (жёлтые ячейки).",
                           "Справа — шаги: где формула, формула в русском Excel и правильный ответ.",
                           "Шаги «измени …» ученик выполняет сам: в файле данные исходные, а ответ к такому шагу — после изменения.",
                           f"Формулы записаны так, что русский Excel покажет их по-русски: {CFG['fns']}."],
                          [f"«Excel-де жұмыс» кезеңінің кілті («{CFG['title']}» сабағы, 8 сынып).",
                           "S1–S5 парақтарында — сценарий кестесі және оқушылардың барлық формуласы жазылған (сары ұяшықтар).",
                           "Оң жақта — қадамдар: формула қайда, орыс тіліндегі Excel-дегі формула және дұрыс жауап.",
                           "«Өзгерт …» қадамдарын оқушы өзі орындайды: файлда бастапқы деректер, ал мұндай қадамның жауабы — өзгерткеннен кейін.",
                           f"Формулалар орыс тіліндегі Excel оларды орысша көрсететіндей жазылған: {CFG['fns']}. Функция атаулары Excel-дің тіліне байланысты көрсетіледі."]), 1):
        info.cell(i, 1, r)
    info.column_dimensions["A"].width = 110
    for sc in SC:
        ws = wb.create_sheet(sc["id"]); put_data(ws, sc); solve(ws, sc, changes=False)
        for st in steps(sc):
            for k in st["c"]: ws[k].fill = YEL
        for j, h in enumerate(tr(["Шаг", "Тип шага", "Где", "Формула (русский Excel)", "Ответ"], ["Қадам", "Қадам түрі", "Қайда", "Формула (орыс тіліндегі Excel)", "Жауап"])):
            c = ws.cell(1, C0 + j, h); c.font = WHITE; c.fill = HEAD
        for i, t in enumerate(sc["tasks"], 1):
            where = tr(f"{t['fill']} (ответ в {t['cell']})", f"{t['fill']} (жауап {t['cell']} ұяшығында)") if t.get("fill") else t["cell"]
            f = (", ".join(clear_or_change(k, v, True) for k, v in t["change"].items())) if t.get("change") else t["f"] + (tr(" и протянуть вниз", " және төмен созып көшіру") if t.get("fill") else "")
            for j, v in enumerate([i, t["tag"], where, f, t["ans"]]):
                c = ws.cell(i + 1, C0 + j, v)
                if isinstance(v, str) and v.startswith("="): c.data_type = "s"   # текст формулы, а не формула
                c.border = Border(bottom=thin); c.alignment = Alignment(wrap_text=True, vertical="top")
        for c, w in {C0: 6, C0 + 1: 24, C0 + 2: 16, C0 + 3: 48, C0 + 4: 12}.items(): ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(path); print("saved", path)

if MODE:
    selfcheck(ARGS[2], MODE)
else:
    selfcheck(LESSON / CFG["check"])
    key(LESSON / CFG["key"])
